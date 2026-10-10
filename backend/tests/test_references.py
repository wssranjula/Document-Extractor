from io import BytesIO

from docx import Document

from app.db import SessionLocal
from app.models import Reference
from tests.test_authz import _signup


def _formulary_docx() -> bytes:
    document = Document()
    document.add_heading("Example Formulary", level=1)
    document.add_heading("Velantine", level=2)
    document.add_paragraph("Standard adult dose: 10 mg orally once daily.")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_authenticated_user_can_upload_a_reference(client):
    token = _signup(client, "reference-owner@example.com")

    response = client.post(
        "/references",
        headers={"Authorization": f"Bearer {token}"},
        data={"name": "  Example   Formulary  "},
        files={
            "file": (
                "example.docx",
                _formulary_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 201
    body = response.json()
    assert body["name"] == "Example Formulary"
    assert body["filename"] == "example.docx"
    assert body["indexed"] is False
    with SessionLocal() as session:
        stored = session.get(Reference, body["id"])
        assert stored is not None
        assert stored.storage_path.endswith(".docx")


def test_reference_upload_requires_authentication_and_docx(client):
    unauthenticated = client.post(
        "/references",
        data={"name": "Formulary"},
        files={"file": ("reference.docx", _formulary_docx(), "application/octet-stream")},
    )
    assert unauthenticated.status_code == 401

    token = _signup(client, "reference-validation@example.com")
    wrong_type = client.post(
        "/references",
        headers={"Authorization": f"Bearer {token}"},
        data={"name": "Formulary"},
        files={"file": ("reference.pdf", b"%PDF", "application/pdf")},
    )
    assert wrong_type.status_code == 422
    assert wrong_type.json()["detail"] == "Upload a DOCX formulary"


def test_reference_requires_heading_two_monographs(client):
    token = _signup(client, "reference-structure@example.com")
    document = Document()
    document.add_heading("Unstructured Formulary", level=1)
    document.add_paragraph("Velantine: 10 mg once daily.")
    stream = BytesIO()
    document.save(stream)

    response = client.post(
        "/references",
        headers={"Authorization": f"Bearer {token}"},
        data={"name": "Unstructured"},
        files={"file": ("unstructured.docx", stream.getvalue(), "application/octet-stream")},
    )

    assert response.status_code == 422
    assert "Heading 2" in response.json()["detail"]
